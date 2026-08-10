"""Pipeline: Excel (optional) -> files/*.txt -> bucr.zip -> api/*.json.

Usage:
    uv run build.py                            # files/*.txt -> bucr.zip + api/*.json
    uv run build.py --from-excel horario.xlsx  # regenerate files/*.txt from Excel first
"""

import argparse
import csv
import datetime
import json
import os
import zipfile

import openpyxl

SPEC = [
    "agency",
    "stops",
    "routes",
    "trips",
    "stop_times",
    "calendar",
    "calendar_dates",
    "fare_attributes",
    "fare_rules",
    "shapes",
    "feed_info",
    "translations",
]

FORBIDDEN_ROUTE = "bUCR_L2"


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime("%H:%M:%S")
        return v.strftime("%Y%m%d")
    if isinstance(v, datetime.date):
        return v.strftime("%Y%m%d")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    s = str(v).strip()
    if s.upper() in ("#NAME?", "NONE", "NAN"):
        return ""
    return s


def export_txt_from_excel(xlsx_path, files_dir):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    for sheet_name in SPEC:
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Missing sheet '{sheet_name}' in {xlsx_path}")
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        # drop phantom columns (empty header)
        cols = [(i, h) for i, h in enumerate(header) if h not in (None, "")]
        idxs = [i for i, _ in cols]
        names = [h for _, h in cols]

        rows = []
        for r in range(2, ws.max_row + 1):
            rv = [ws.cell(row=r, column=i + 1).value for i in idxs]
            if rv[0] in (None, ""):
                continue  # blank filler row
            rows.append(rv)

        if sheet_name == "stops" and "stop_id" in names:
            si = names.index("stop_id")
            if "stop_url" in names:
                j = names.index("stop_url")
                for rv in rows:
                    rv[j] = f"https://bus.ucr.ac.cr/paradas/{rv[si]}"
            if "stop_point" in names and "stop_lat" in names and "stop_lon" in names:
                j = names.index("stop_point")
                lat_i = names.index("stop_lat")
                lon_i = names.index("stop_lon")
                for rv in rows:
                    lat = str(rv[lat_i]).strip()
                    lon = str(rv[lon_i]).strip()
                    rv[j] = f"SRID=4326;POINT ({lon} {lat})"

        out_path = os.path.join(files_dir, sheet_name + ".txt")
        with open(out_path, "w", newline="\n", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(names)
            for rv in rows:
                w.writerow([fmt(v) for v in rv])


def read_txt(files_dir, name):
    path = os.path.join(files_dir, name + ".txt")
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def validate(files_dir):
    tables = {}
    for name in SPEC:
        path = os.path.join(files_dir, name + ".txt")
        if os.path.exists(path):
            tables[name] = read_txt(files_dir, name)

    def col(table, field):
        return {r[field] for r in tables[table]}

    required = ["routes", "trips", "stop_times", "stops", "calendar", "shapes"]
    missing = [n for n in required if n not in tables]
    if missing:
        raise SystemExit(f"Missing required GTFS files: {missing}")

    errors = []

    if not col("trips", "route_id") <= col("routes", "route_id"):
        errors.append("trips.route_id has values not present in routes.route_id")
    if not col("trips", "service_id") <= col("calendar", "service_id"):
        errors.append("trips.service_id has values not present in calendar.service_id")
    if not col("trips", "shape_id") <= col("shapes", "shape_id"):
        errors.append("trips.shape_id has values not present in shapes.shape_id")
    if not col("stop_times", "trip_id") <= col("trips", "trip_id"):
        errors.append("stop_times.trip_id has values not present in trips.trip_id")
    if not col("trips", "trip_id") <= col("stop_times", "trip_id"):
        errors.append("there are trips with no stop_times")
    if not col("stop_times", "stop_id") <= col("stops", "stop_id"):
        errors.append("stop_times.stop_id has values not present in stops.stop_id")
    if "fare_rules" in tables and tables["fare_rules"]:
        if not col("fare_rules", "route_id") <= col("routes", "route_id"):
            errors.append(
                "fare_rules.route_id has values not present in routes.route_id"
            )
    if "calendar_dates" in tables and tables["calendar_dates"]:
        if not col("calendar_dates", "service_id") <= col("calendar", "service_id"):
            errors.append(
                "calendar_dates.service_id has values not present in calendar.service_id"
            )

    from collections import defaultdict

    by_trip = defaultdict(list)
    for r in tables["stop_times"]:
        by_trip[r["trip_id"]].append(int(r["stop_sequence"]))
    unordered = [t for t, seq in by_trip.items() if seq != sorted(seq)]
    if unordered:
        errors.append(
            f"unordered stop_sequence in {len(unordered)} trip(s): {unordered[:5]}"
        )

    if len(tables["routes"]) != 1:
        errors.append(f"expected exactly 1 route, found {len(tables['routes'])}")

    for name in SPEC:
        path = os.path.join(files_dir, name + ".txt")
        if os.path.exists(path):
            content = open(path, encoding="utf-8").read()
            if FORBIDDEN_ROUTE in content:
                errors.append(f"'{FORBIDDEN_ROUTE}' found in {name}.txt")
            if "#NAME?" in content:
                errors.append(f"'#NAME?' found in {name}.txt")

    if errors:
        raise SystemExit(
            "Referential integrity check failed:\n  - " + "\n  - ".join(errors)
        )

    print(
        f"Referential integrity OK ({len(tables)} files, {len(tables['trips'])} trips)"
    )


def build_zip(files_dir, zip_path):
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in SPEC:
            path = os.path.join(files_dir, name + ".txt")
            if os.path.exists(path):
                z.write(path, arcname=name + ".txt")


def build_json(files_dir, api_dir):
    os.makedirs(api_dir, exist_ok=True)
    index = {}
    for name in SPEC:
        path = os.path.join(files_dir, name + ".txt")
        if not os.path.exists(path):
            continue
        rows = read_txt(files_dir, name)
        with open(os.path.join(api_dir, name + ".json"), "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=2)
            f.write("\n")
        index[name] = len(rows)
    with open(os.path.join(api_dir, "index.json"), "w", encoding="utf-8") as f:
        json.dump(
            {
                "files": index,
                "generated_at": datetime.datetime.now(datetime.timezone.utc)
                .isoformat()
                .replace("+00:00", "Z"),
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
        f.write("\n")


def build_geojson(files_dir, api_dir):
    """shapes.geojson / stops.geojson, following the same convention as
    incofer's utils/create_geo_shapes.py and create_geo_stops.py: GeoJSON is
    generated once here at build time, not re-derived by every consumer."""
    os.makedirs(api_dir, exist_ok=True)

    shapes = read_txt(files_dir, "shapes")
    by_shape = {}
    for row in shapes:
        by_shape.setdefault(row["shape_id"], []).append(row)

    shape_features = []
    for shape_id, points in by_shape.items():
        points.sort(key=lambda p: int(p["shape_pt_sequence"]))
        coordinates = [
            [float(p["shape_pt_lon"]), float(p["shape_pt_lat"])] for p in points
        ]
        last_point = points[-1]
        shape_features.append(
            {
                "type": "Feature",
                "geometry": {"type": "LineString", "coordinates": coordinates},
                "properties": {
                    "shape_id": shape_id,
                    "shape_dist_traveled": float(last_point["shape_dist_traveled"])
                    if last_point.get("shape_dist_traveled")
                    else None,
                },
            }
        )
    with open(os.path.join(api_dir, "shapes.geojson"), "w", encoding="utf-8") as f:
        json.dump(
            {"type": "FeatureCollection", "features": shape_features},
            f,
            ensure_ascii=False,
            indent=2,
        )
        f.write("\n")

    stops = read_txt(files_dir, "stops")
    stop_features = []
    for row in stops:
        props = dict(row)
        lon = float(props.pop("stop_lon"))
        lat = float(props.pop("stop_lat"))
        props.pop("stop_point", None)  # redundant WKT encoding of the same coordinates
        stop_features.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [lon, lat]},
                "properties": props,
            }
        )
    with open(os.path.join(api_dir, "stops.geojson"), "w", encoding="utf-8") as f:
        json.dump(
            {"type": "FeatureCollection", "features": stop_features},
            f,
            ensure_ascii=False,
            indent=2,
        )
        f.write("\n")


def main():
    root = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--from-excel",
        help="Path to the master Excel workbook (.xlsx/.xlsm) to regenerate files/*.txt from",
    )
    args = ap.parse_args()

    files_dir = os.path.join(root, "files")
    os.makedirs(files_dir, exist_ok=True)

    if args.from_excel:
        export_txt_from_excel(args.from_excel, files_dir)
        print(f"files/*.txt regenerated from {args.from_excel}")

    validate(files_dir)
    build_zip(files_dir, os.path.join(root, "bucr.zip"))
    build_json(files_dir, os.path.join(root, "api"))
    build_geojson(files_dir, os.path.join(root, "api"))
    print("build OK")


if __name__ == "__main__":
    main()
